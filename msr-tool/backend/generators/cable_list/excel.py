"""Generator: Kabelzugliste einer Anlage als Excel (.xlsx).

Ableitung (siehe /docs/msr-tool/ER-Modell.md §5): nur Hardware-Datenpunkte
(dp_objekttyp.ist_hardware) werden verdrahtet. Signale werden je Feldgerät
gruppiert – das Feldgerät ist der BAS bis einschließlich Betriebsmittel-Block
(Funktionsblock abgeschnitten). Je Feldgerät ein Kabel (fortlaufende Kabel-Nr);
Kabeltyp/Adern/Querschnitt aus der Klemmenvorlage (objektspezifisch, sonst je
Signalart als Default).
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_HW_QUERY = """
SELECT d.bas, d.bas_funktion, o.uuid, o.object_type, o.bezeichnung_de
FROM project.datenpunkt d
JOIN catalog.dp_objekttyp o ON o.uuid = d.dp_objekttyp_uuid
WHERE d.anlage_id = %s AND o.ist_hardware
ORDER BY d.bas
"""

_VORLAGE_QUERY = """
SELECT kv.dp_objekttyp_uuid, kv.signalart, kv.klemme, kv.adern, kv.querschnitt,
       kt.bezeichnung AS kabeltyp
FROM catalog.klemmen_vorlage kv
LEFT JOIN catalog.kabeltyp kt ON kt.id = kv.kabeltyp_id
"""

_HEADERS = [
    ("Kabel-Nr.", 10), ("Feldgerät (BAS)", 40), ("Von", 16), ("Nach", 16),
    ("Signal", 10), ("Beschreibung", 32), ("Signalart", 10), ("Klemme", 10),
    ("Adern", 7), ("Kabeltyp", 22), ("Querschnitt", 11),
]


def _device_of(bas: str, funktion: str | None) -> str:
    """Feldgerät = BAS ohne den Funktionsblock (letztes Segment)."""
    if funktion and bas.endswith("_" + funktion):
        return bas[: -(len(funktion) + 1)]
    return bas.rsplit("_", 1)[0]


def _load_vorlagen(cur):
    cur.execute(_VORLAGE_QUERY)
    by_uuid, by_signalart = {}, {}
    for uuid, signalart, klemme, adern, querschnitt, kabeltyp in cur.fetchall():
        spec = {"klemme": klemme, "adern": adern,
                "querschnitt": querschnitt, "kabeltyp": kabeltyp}
        if uuid is not None:
            by_uuid[uuid] = spec
        elif signalart is not None:
            by_signalart[signalart] = spec
    return by_uuid, by_signalart


def export_cable_list(cur, anlage_id: int, path: str, nach: str = "Schaltschrank") -> int:
    """Schreibt die Kabelzugliste der Anlage nach `path`. Rückgabe: Signal-Zeilen."""
    by_uuid, by_signalart = _load_vorlagen(cur)
    cur.execute(_HW_QUERY, (anlage_id,))
    rows = cur.fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Kabelzugliste"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for col, (title, width) in enumerate(_HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=title)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = width

    kabel_nr: dict[str, str] = {}
    r = 2
    for bas, funktion, uuid, object_type, bez in rows:
        device = _device_of(bas, funktion)
        if device not in kabel_nr:
            kabel_nr[device] = f"W{len(kabel_nr) + 1:02d}"
        spec = by_uuid.get(uuid) or by_signalart.get(object_type) or {}

        ws.cell(row=r, column=1, value=kabel_nr[device])
        ws.cell(row=r, column=2, value=device)
        ws.cell(row=r, column=3, value=device)
        ws.cell(row=r, column=4, value=nach)
        ws.cell(row=r, column=5, value=funktion)
        ws.cell(row=r, column=6, value=bez)
        ws.cell(row=r, column=7, value=object_type)
        ws.cell(row=r, column=8, value=spec.get("klemme"))
        ws.cell(row=r, column=9, value=spec.get("adern"))
        ws.cell(row=r, column=10, value=spec.get("kabeltyp"))
        ws.cell(row=r, column=11, value=spec.get("querschnitt"))
        r += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(_HEADERS)).column_letter}{len(rows) + 1}"
    wb.properties.title = f"Kabelzugliste ({len(kabel_nr)} Kabel)"
    wb.save(path)
    return len(rows)
