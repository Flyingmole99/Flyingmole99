"""Live-API-Test (FastAPI TestClient). Nur mit BACTWIN_DIR + DATABASE_URL.

Spielt den Drag-&-Drop-Kernablauf durch: Palette -> Projekt -> Anlage ->
Baugruppe ziehen (instanziiert) -> Datenpunkte -> Dokument-Downloads.
"""
import io
import os

import pytest

BACTWIN_DIR = os.environ.get("BACTWIN_DIR")
DATABASE_URL = os.environ.get("DATABASE_URL")

pytestmark = pytest.mark.skipif(
    not (BACTWIN_DIR and DATABASE_URL),
    reason="BACTWIN_DIR und DATABASE_URL erforderlich",
)


def _cleanup():
    import psycopg2
    conn = psycopg2.connect(DATABASE_URL)
    with conn, conn.cursor() as cur:
        cur.execute("DELETE FROM project.projekt WHERE name='PYTEST_API'")
    conn.close()


def test_drag_and_drop_flow():
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook

    from api.app import app
    from importer.run_import import run

    run(BACTWIN_DIR, DATABASE_URL, "AMEV 1.2 (test)")
    _cleanup()

    with TestClient(app) as client:
        # Palette: Kessel-Template finden
        r = client.get("/catalog/aggregate-templates", params={"q": "BGP_KES_nM"})
        assert r.status_code == 200
        tmpl = next(t for t in r.json() if t["kennung"] == "BGP_KES_nM_AMEV1")

        # Projekt
        r = client.post("/projekte", json={"name": "PYTEST_API", "kunde": "Test"})
        assert r.status_code == 201, r.text
        projekt_id = r.json()["id"]

        # Anlage
        r = client.post(f"/projekte/{projekt_id}/anlagen",
                        json={"gewerk_kg": "420", "anlage_kuerzel": "EZA", "nummer": 1})
        assert r.status_code == 201, r.text
        anlage = r.json()
        anlage_id = anlage["id"]
        assert anlage["bas"] == "420_EZA01"

        # Drag & Drop: Kessel auf die Anlage ziehen -> instanziiert
        r = client.post(f"/anlagen/{anlage_id}/baugruppen",
                        json={"aggregat_template_uuid": tmpl["uuid"]})
        assert r.status_code == 201, r.text
        bg = r.json()
        assert bg["kuerzel"] == "KES"
        assert bg["bas"] == "420_EZA01_KES01"
        assert bg["instanziierung"]["hardware"] == 22

        # Datenpunkte
        r = client.get(f"/anlagen/{anlage_id}/datenpunkte")
        assert r.status_code == 200
        dps = r.json()
        assert len(dps) == bg["instanziierung"]["datenpunkte"]
        assert all(d["bas"].startswith("420_EZA01_KES01_") for d in dps)
        assert sum(1 for d in dps if d["hardware"]) == 22

        # Dokument-Downloads
        for pfad, sheet in [("datenpunktliste", "Datenpunktliste"),
                            ("kabelzugliste", "Kabelzugliste")]:
            r = client.get(f"/anlagen/{anlage_id}/dokumente/{pfad}")
            assert r.status_code == 200, r.text
            assert "spreadsheetml" in r.headers["content-type"]
            wb = load_workbook(io.BytesIO(r.content))
            assert wb.active.max_row > 1

        # Anlage-Ansicht zeigt die Baugruppe
        r = client.get(f"/anlagen/{anlage_id}")
        assert r.status_code == 200
        assert r.json()["baugruppen"][0]["datenpunkte"] == bg["instanziierung"]["datenpunkte"]

    _cleanup()
