"""Live-Test: Import -> Baugruppe instanziieren -> Datenpunktliste exportieren.

Läuft nur mit BACTWIN_DIR + DATABASE_URL (DB mit angewandten Migrationen).
"""
import os
import tempfile

import pytest

BACTWIN_DIR = os.environ.get("BACTWIN_DIR")
DATABASE_URL = os.environ.get("DATABASE_URL")

pytestmark = pytest.mark.skipif(
    not (BACTWIN_DIR and DATABASE_URL),
    reason="BACTWIN_DIR und DATABASE_URL erforderlich",
)


def _reset_project(cur):
    cur.execute("DELETE FROM project.projekt WHERE name = 'PYTEST_KESSEL'")


def test_instantiate_and_export():
    import psycopg2
    from openpyxl import load_workbook

    from domain.instantiate import instantiate_baugruppe
    from generators.datapoint_list.excel import export_datapoint_list
    from importer.run_import import run

    run(BACTWIN_DIR, DATABASE_URL, "AMEV 1.2 (test)")

    conn = psycopg2.connect(DATABASE_URL)
    conn.set_client_encoding("UTF8")
    try:
        with conn, conn.cursor() as cur:
            _reset_project(cur)
            # Projektgraph anlegen
            cur.execute("SELECT id FROM catalog.naming_profile WHERE variante='ungekuerzt' LIMIT 1")
            profile_id = cur.fetchone()[0]
            cur.execute("SELECT id FROM catalog.bac_version ORDER BY id DESC LIMIT 1")
            version_id = cur.fetchone()[0]
            cur.execute("""INSERT INTO project.projekt(name, naming_profile_id, bac_version_id)
                           VALUES ('PYTEST_KESSEL', %s, %s) RETURNING id""",
                        (profile_id, version_id))
            projekt_id = cur.fetchone()[0]
            cur.execute("""INSERT INTO project.anlage(projekt_id, gewerk_kg, anlage_kuerzel, nummer, bas)
                           VALUES (%s,'420','EZA',1,'420_EZA01') RETURNING id""", (projekt_id,))
            anlage_id = cur.fetchone()[0]
            cur.execute("SELECT uuid FROM catalog.aggregat_template WHERE kennung='BGP_KES_nM_AMEV1'")
            tmpl = cur.fetchone()[0]
            cur.execute("""INSERT INTO project.baugruppe(anlage_id, aggregat_template_uuid, kuerzel, nummer, medium_pos, bas)
                           VALUES (%s,%s,'KES',1,'HZ~','420_EZA01_KES01') RETURNING id""",
                        (anlage_id, tmpl))
            baugruppe_id = cur.fetchone()[0]

            stats = instantiate_baugruppe(cur, baugruppe_id)

            # Akzeptanz: aus der bekannten Roh-Expansion (76 Objekte / 22 Hardware)
            assert stats["objekte_gesamt"] == 76, stats
            assert stats["container_uebersprungen"] == 6, stats   # SV-Container
            assert stats["hardware"] == 22, stats                 # AI/AO/BI/BO
            assert stats["trend"] > 0, stats                      # TL gefaltet
            assert stats["varianten_dubletten"] > 0, stats        # BI_HD/MI_HD etc.
            assert stats["datenpunkte"] < 76, stats

            # BAS ist je Anlage eindeutig und beginnt mit dem Baugruppen-Kontext
            cur.execute("SELECT count(*), count(DISTINCT bas) FROM project.datenpunkt WHERE anlage_id=%s", (anlage_id,))
            n_total, n_distinct = cur.fetchone()
            assert n_total == n_distinct == stats["datenpunkte"], (n_total, n_distinct)
            cur.execute("SELECT count(*) FROM project.datenpunkt WHERE anlage_id=%s AND bas NOT LIKE '420\\_EZA01\\_KES01\\_%%'", (anlage_id,))
            assert cur.fetchone()[0] == 0

            # Export Datenpunktliste
            with tempfile.TemporaryDirectory() as td:
                path = os.path.join(td, "dp.xlsx")
                n = export_datapoint_list(cur, anlage_id, path)
                assert n == stats["datenpunkte"]
                wb = load_workbook(path)
                ws = wb.active
                assert ws.max_row == n + 1          # + Kopfzeile
                assert ws.cell(row=1, column=1).value == "Datenpunkt (BAS)"

            # Export Kabelzugliste: eine Zeile je Hardware-Signal
            from generators.cable_list.excel import export_cable_list
            with tempfile.TemporaryDirectory() as td:
                path = os.path.join(td, "kabel.xlsx")
                n = export_cable_list(cur, anlage_id, path)
                assert n == stats["hardware"] == 22
                wb = load_workbook(path)
                ws = wb.active
                assert ws.cell(row=1, column=1).value == "Kabel-Nr."
                # jede Signalzeile hat einen Kabeltyp und eine Kabel-Nr.
                for row in ws.iter_rows(min_row=2, values_only=True):
                    assert row[0] and row[0].startswith("W")     # Kabel-Nr.
                    assert row[9]                                  # Kabeltyp
                # mehrere Feldgeräte -> mehrere distinct Kabel
                kabel = {row[0] for row in ws.iter_rows(min_row=2, values_only=True)}
                assert 1 < len(kabel) <= 22

            _reset_project(cur)
    finally:
        conn.close()
